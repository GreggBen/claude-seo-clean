import contextlib
import importlib.util
import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook


SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "seo" / "scripts"


def load_script(name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    with patch.object(sys, "path", [str(SCRIPTS), *sys.path]):
        spec.loader.exec_module(module)
    return module


gsc_query = load_script("gsc_query")
# Les tests HTML/XLSX ne chargent pas les bibliothèques natives du moteur PDF.
with patch.dict(sys.modules, {"weasyprint": Mock(HTML=Mock(side_effect=AssertionError("PDF rendering is outside these tests")))}):
    google_report = load_script("google_report")


class GscReportingTests(unittest.TestCase):
    def query(self, responses):
        service = Mock()
        service.searchanalytics.return_value.query.return_value.execute.side_effect = responses
        with patch.object(gsc_query, "_build_gsc_service", return_value=service):
            return gsc_query.query_search_analytics(
                "sc-domain:example.com", "2026-01-01", "2026-01-28",
                dimensions=["query"],
            )

    def assert_unavailable_metrics(self, result):
        self.assertEqual(result["totals"], dict.fromkeys(["clicks", "impressions", "ctr", "position"]))

    def test_missing_service_is_error_without_zero_totals(self):
        with patch.object(gsc_query, "_build_gsc_service", return_value=None):
            result = gsc_query.query_search_analytics("sc-domain:example.com")
        self.assert_unavailable_metrics(result)
        self.assertEqual(result["status"], "ERROR")
        self.assertEqual(result["totals_status"], "ERROR")
        self.assertIsNone(result["row_count"])

    def test_api_error_is_not_zero_traffic(self):
        result = self.query([RuntimeError("403 permission denied")])
        self.assert_unavailable_metrics(result)
        self.assertEqual(result["status"], "ERROR")
        self.assertIsNone(result["totals_source"])
        self.assertIn("Permission denied", result["error"])

    def test_empty_response_marks_data_unavailable(self):
        result = self.query([{}, {}])
        self.assert_unavailable_metrics(result)
        self.assertEqual(result["status"], "NOT_AVAILABLE")
        self.assertEqual(result["totals_status"], "NOT_AVAILABLE")
        self.assertEqual(result["row_count"], 0)
        self.assertEqual(result["totals_source"], "dimensionless_aggregate")

    def test_explicit_zero_counts_remain_measured(self):
        result = self.query([{}, {"rows": [{"clicks": 0, "impressions": 0, "ctr": 0, "position": 0}]}])
        self.assertEqual(result["status"], "OK")
        self.assertEqual(result["totals"]["clicks"], 0)
        self.assertEqual(result["totals"]["impressions"], 0)
        self.assertIsNone(result["totals"]["position"])
        self.assertIsNone(result["totals"]["ctr"])

    def test_aggregate_totals_preserve_anonymized_traffic(self):
        result = self.query([
            {"rows": [{"keys": ["query"], "clicks": 5, "impressions": 100, "ctr": 0.05, "position": 2}]},
            {"rows": [{"clicks": 10, "impressions": 200, "ctr": 0.05, "position": 3}]},
        ])
        self.assertEqual(result["status"], "OK")
        self.assertEqual(result["totals"]["clicks"], 10)
        self.assertEqual(result["totals_source"], "dimensionless_aggregate")
        self.assertEqual(result["ctr_unit"], "percent")
        self.assertEqual(result["rows"][0]["ctr"], 5)

    def test_row_sum_fallback_exposes_partial_coverage_and_failure(self):
        result = self.query([
            {"rows": [{"keys": ["query"], "clicks": 5, "impressions": 100, "ctr": 0.05, "position": 2}]},
            RuntimeError("aggregate unavailable"),
        ])
        self.assertEqual(result["status"], "PARTIAL")
        self.assertEqual(result["totals_status"], "PARTIAL")
        self.assertEqual(result["totals_source"], "dimension_row_sum")
        self.assertIn("aggregate unavailable", result["totals_error"])
        self.assertTrue(result["limits"])
        self.assertEqual(result["totals"]["clicks"], 5)
        self.assertEqual(result["totals"]["ctr"], 5)
        self.assertIsNone(result["totals"]["position"])

    def test_empty_rows_do_not_provide_fallback_zero_totals(self):
        result = self.query([{}, RuntimeError("aggregate unavailable")])
        self.assert_unavailable_metrics(result)
        self.assertEqual(result["status"], "ERROR")
        self.assertEqual(result["totals_status"], "ERROR")

    def test_missing_metrics_remain_null(self):
        result = self.query([{"rows": [{"keys": ["query"], "impressions": 100}]}, {"rows": [{"impressions": 100}]}])
        self.assertIsNone(result["totals"]["clicks"])
        self.assertIsNone(result["totals"]["ctr"])
        self.assertIsNone(result["rows"][0]["position"])
        self.assertEqual(result["totals_status"], "PARTIAL")

    def export_ctr(self, gsc, sheet="Queries"):
        with tempfile.TemporaryDirectory() as output_dir:
            path = google_report.generate_xlsx({"gsc": gsc}, "example.com", "gsc-performance", output_dir)
            workbook = load_workbook(path)
            try:
                return workbook[sheet]["D2"].value
            finally:
                workbook.close()

    def test_producer_ctr_is_five_percent_in_both_xlsx_sheets(self):
        result = self.query([
            {"rows": [{"keys": ["query"], "clicks": 5, "impressions": 100, "ctr": 0.05, "position": 2}]},
            {"rows": [{"clicks": 5, "impressions": 100, "ctr": 0.05, "position": 2}]},
        ])
        result["pages"] = result["rows"]
        self.assertEqual(self.export_ctr(result), "5.00%")
        self.assertEqual(self.export_ctr(result, "Pages"), "5.00%")

    def test_legacy_processed_percentage_remains_supported(self):
        self.assertEqual(self.export_ctr({"rows": [{"query": "query", "ctr": 5}]}), "5.00%")

    def test_explicit_ratio_is_supported(self):
        self.assertEqual(self.export_ctr({"ctr_unit": "ratio", "rows": [{"ctr": 0.05}]}), "5.00%")

    def test_xlsx_missing_ctr_is_unavailable_not_zero(self):
        self.assertEqual(self.export_ctr({"rows": [{"query": "query"}]}), "NOT_AVAILABLE")

    def test_measured_zero_ctr_is_exported_as_zero_percent(self):
        result = self.query([
            {"rows": [{"keys": ["query"], "clicks": 0, "impressions": 100, "ctr": 0, "position": 2}]},
            {"rows": [{"clicks": 0, "impressions": 100, "ctr": 0, "position": 2}]},
        ])
        self.assertEqual(result["totals"]["ctr"], 0)
        self.assertEqual(self.export_ctr(result), "0.00%")

    def test_raw_producer_output_is_exportable_without_wrapper(self):
        with tempfile.TemporaryDirectory() as output_dir:
            path = google_report.generate_xlsx({"rows": [{"query": "query", "ctr": 5}]}, "example.com", "gsc-performance", output_dir)
            workbook = load_workbook(path)
            try:
                self.assertEqual(workbook["Queries"]["D2"].value, "5.00%")
            finally:
                workbook.close()

    def test_html_handles_null_totals_and_shows_failure(self):
        data = {"status": "ERROR", "totals_status": "ERROR", "error": "Permission denied", "totals": dict.fromkeys(["clicks", "impressions", "ctr", "position"]), "row_count": None}
        section, _ = google_report._build_gsc_section(data, {})
        self.assertIn("ERROR", section)
        self.assertIn("Permission denied", section)
        with tempfile.TemporaryDirectory() as output_dir:
            report = google_report.generate_report("gsc-performance", data, "example.com", output_dir, "html")
            self.assertIsNone(report["error"])
            html_path = next(Path(path) for path in report["files"] if str(path).endswith(".html"))
            self.assertIn("NOT_AVAILABLE", html_path.read_text())

    def test_html_exposes_fallback_provenance(self):
        data = self.query([{"rows": [{"keys": ["query"], "clicks": 5, "impressions": 100, "ctr": 0.05, "position": 2}]}, RuntimeError("aggregate unavailable")])
        section, _ = google_report._build_gsc_section(data, {})
        self.assertIn("dimension_row_sum", section)
        self.assertIn("PARTIAL", section)
        self.assertIn("aggregate unavailable", section)

    def test_html_and_chart_accept_missing_row_metrics(self):
        data = self.query([{"rows": [{"keys": ["query"], "impressions": 100}]}, {"rows": [{"impressions": 100}]}])
        with tempfile.TemporaryDirectory() as output_dir:
            plot = Mock()
            axis = Mock()
            plot.subplots.return_value = (Mock(), axis)
            axis.barh.return_value = [Mock()]
            with patch.object(google_report, "plt", plot):
                google_report.chart_top_queries(data, Path(output_dir))
            self.assertEqual(axis.barh.call_count, 1)
            self.assertEqual(axis.barh.call_args.args[1], [100])
            with patch.object(google_report, "chart_top_queries", return_value=""):
                report = google_report.generate_report("gsc-performance", {"gsc": data}, "example.com", output_dir, "html")
            self.assertIsNone(report["error"])
            html_path = next(Path(path) for path in report["files"] if str(path).endswith(".html"))
            content = html_path.read_text()
            self.assertIn("PARTIAL", content)
            self.assertIn("NOT_AVAILABLE", content)
            self.assertNotIn("Queries in Top 3", content)

    def test_json_cli_error_has_failure_exit_code(self):
        with patch.object(sys, "argv", ["gsc_query.py", "--property", "sc-domain:example.com", "--json"]), patch.object(gsc_query, "_build_gsc_service", return_value=None):
            stdout = io.StringIO()
            with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as raised:
                gsc_query.main()
        self.assertEqual(raised.exception.code, 1)
        self.assertIn('"clicks": null', stdout.getvalue())

    def test_empty_cli_prints_named_absence(self):
        with patch.object(sys, "argv", ["gsc_query.py", "--property", "sc-domain:example.com"]), patch.object(gsc_query, "_build_gsc_service") as factory:
            factory.return_value.searchanalytics.return_value.query.return_value.execute.side_effect = [{}, {}]
            stdout = io.StringIO()
            with contextlib.redirect_stdout(stdout):
                gsc_query.main()
        self.assertIn("NOT_AVAILABLE", stdout.getvalue())


if __name__ == "__main__":
    unittest.main()
